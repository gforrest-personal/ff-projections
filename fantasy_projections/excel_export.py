"""Write the draft-day Excel workbook, the tool's one output meant for people.

fantasy_draft_board.xlsx has an ALL sheet (the full draft board) plus one sheet
per position. Position sheets add a `drop_to_next` column — the average_points a
player projects above the next player AT THE SAME POSITION — so you can see
positional cliffs on draft day. This module only formats and writes; it never
touches collection or matching.
"""

from __future__ import annotations

import pandas as pd
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter

from fantasy_projections.config import XLSX_FILE, POSITION_LIMITS

# Point columns that should render with two decimals.
POINT_COLS = {
    "sleeper_points",
    "espn_points",
    "yahoo_points",
    "average_points",
    "drop_to_next",
}

ALL_COLUMNS = [
    "player",
    "team",
    "position",
    "sleeper_points",
    "espn_points",
    "yahoo_points",
    "average_points",
]

POSITION_COLUMNS = [
    "player",
    "team",
    "sleeper_points",
    "espn_points",
    "yahoo_points",
    "average_points",
    "drop_to_next",
]

# Reasonable on-screen widths per column name.
COLUMN_WIDTHS = {
    "player": 24,
    "team": 7,
    "position": 10,
    "sleeper_points": 15,
    "espn_points": 14,
    "yahoo_points": 14,
    "average_points": 16,
    "drop_to_next": 14,
}

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _all_sheet(board: pd.DataFrame) -> pd.DataFrame:
    df = board.sort_values("average_points", ascending=False)
    return df[ALL_COLUMNS].reset_index(drop=True)


def _position_sheet(board: pd.DataFrame, position: str) -> pd.DataFrame:
    """Filter to one position, sort by average, and add drop_to_next.

    drop_to_next = this player's average_points - the next player's, computed
    strictly within this position. The last player is left blank.
    """
    df = board[board["position"] == position].copy()
    df = df.sort_values("average_points", ascending=False).reset_index(drop=True)

    # diff(periods=-1) gives current - next, i.e. the points this player projects
    # above the next player at this position. The final row has no "next" player,
    # so it stays NaN (blank in Excel).
    df["drop_to_next"] = df["average_points"].diff(periods=-1)

    return df[POSITION_COLUMNS]


def _format_sheet(ws) -> None:
    """Apply freeze panes, filters, bold headers, widths, and number formats."""
    max_row = ws.max_row
    max_col = ws.max_column
    headers = [ws.cell(row=1, column=c).value for c in range(1, max_col + 1)]

    # Header styling.
    for c in range(1, max_col + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Freeze the header row and enable autofilter over the whole table.
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"

    # Column widths + two-decimal number format on point columns.
    for c, name in enumerate(headers, start=1):
        letter = get_column_letter(c)
        ws.column_dimensions[letter].width = COLUMN_WIDTHS.get(name, 14)
        if name in POINT_COLS:
            for r in range(2, max_row + 1):
                ws.cell(row=r, column=c).number_format = "0.00"

    # Conditional formatting: shade drop_to_next so bigger positional cliffs pop.
    if "drop_to_next" in headers and max_row >= 2:
        col = headers.index("drop_to_next") + 1
        letter = get_column_letter(col)
        rng = f"{letter}2:{letter}{max_row}"
        # White (small drop) -> orange -> red (large drop / positional cliff).
        rule = ColorScaleRule(
            start_type="min", start_color="FFFFFF",
            mid_type="percentile", mid_value=70, mid_color="FFC000",
            end_type="max", end_color="C00000",
        )
        ws.conditional_formatting.add(rng, rule)


def build_workbook(board: pd.DataFrame) -> None:
    """Write fantasy_draft_board.xlsx from the final draft board."""
    sheets = {"ALL": _all_sheet(board)}
    for position in POSITION_LIMITS:  # QB, RB, WR, TE
        sheets[position] = _position_sheet(board, position)

    with pd.ExcelWriter(XLSX_FILE, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name, index=False)
        for name in sheets:
            _format_sheet(writer.sheets[name])

    print(f"Excel: wrote {XLSX_FILE} with sheets {list(sheets)}")
